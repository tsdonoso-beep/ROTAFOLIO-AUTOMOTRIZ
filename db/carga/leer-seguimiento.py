#!/usr/bin/env python3
"""Saca del seguimiento de Control de Gestión lo que la aplicación necesita.

    python3 db/carga/leer-seguimiento.py SEGUIMIENTO.xlsx > /tmp/seguimiento.json

El archivo es un libro de diecisiete hojas que sostiene a mano todo el proceso.
De ahí solo se leen tres, y solo las columnas que tienen destino en la base:

  · «2.1 Seguimiento Memos Personal» — el anexo de cada memo: qué persona,
    cuánto le tocó y entre qué fechas viajó. Es el dato que no existe en
    ninguna otra parte fuera del Word adjunto de cada memo.
  · «4. Estados Contrato Personal Pa» — quiénes son esas personas.
  · «2. Seguimiento Memos Annie» — los memos de hospedaje y caja chica, que
    la hoja 2.1 no cubre.

No escribe nada: imprime JSON y se acabó. Lo que decide qué entra a la base es
`memos-historicos.mts`, y corre en seco salvo que se le pida lo contrario.
"""

import json
import sys
from datetime import date, datetime

import openpyxl


def dni(v):
    """El DNI como lo guarda la base: ocho dígitos, con sus ceros delante.

    Excel guarda los DNI como número y se come el cero inicial, así que
    «3673415» y «03673415» son la misma persona escrita de dos maneras. Sin
    rellenar, uno de cada diez no encuentra a su dueño.
    """
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s.zfill(8) if s.isdigit() else s


def texto(v):
    return "" if v is None else " ".join(str(v).split())


def fecha(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return None


def numero(v):
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return round(n, 2)


def personas(wb):
    ws = wb["4. Estados Contrato Personal Pa"]
    out = {}
    for f in ws.iter_rows(min_row=2, values_only=True):
        d = dni(f[2])
        if not d or d in out:
            continue
        out[d] = {
            "dni": d,
            "nombre": texto(f[3]),
            "cargo": texto(f[4]),
            # CESADO o ACTIVO. Quien cesó y no rindió es plata que no vuelve,
            # así que entra igual: se necesita para saber a quién se le debe.
            "activo": texto(f[8]).upper() == "ACTIVO",
            "banco": texto(f[9]),
            "cuenta": texto(f[10]),
            "cci": texto(f[11]),
            "correo": texto(f[12]).lower(),
        }
    return list(out.values())


def anexo(wb):
    ws = wb["2.1 Seguimiento Memos Personal"]
    out = []
    for f in ws.iter_rows(min_row=4, values_only=True):
        if not f[1] or not f[3]:
            continue
        out.append({
            "correlativo": texto(f[1]),
            "dni": dni(f[3]),
            "nombre": texto(f[4]),
            "proyecto": texto(f[6]),      # viene la abreviatura, no el código
            "tipo": texto(f[7]).upper(),
            "empresa": texto(f[8]),
            "fecha_abono": fecha(f[9]),
            "status": texto(f[10]).upper(),
            "monto": numero(f[11]),
            "rendido": numero(f[12]),
            "facturas": numero(f[13]),
            "dj": numero(f[14]),
            "movilidad": numero(f[15]),
            "devuelto": numero(f[18]),
            "carpeta": texto(f[19]),
            "fecha_desde": fecha(f[22]),
            "fecha_hasta": fecha(f[23]),
        })
    return out


def memos_de_annie(wb):
    """Hospedaje y caja chica, que la hoja del anexo no cubre.

    Esta hoja va por responsable y no por persona, así que de acá salen los
    memos pero no a quiénes cubren. Se cargan igual: un memo sin anexo se ve
    en la aplicación y se completa; un memo que no existe, no.
    """
    ws = wb["2. Seguimiento Memos Annie"]
    out = []
    for f in ws.iter_rows(min_row=4, values_only=True):
        if not f[1]:
            continue
        out.append({
            "correlativo": texto(f[1]),
            "responsable": texto(f[2]),
            "destino": texto(f[3]),
            "proyecto": texto(f[4]),
            "tipo": texto(f[5]).upper(),
            "empresa": texto(f[6]),
            "fecha_abono": fecha(f[7]),
            "monto": numero(f[10]),
            "cuenta": texto(f[11]),
            "status": texto(f[14]).upper(),
        })
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit("Falta la ruta del archivo de seguimiento.")

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=False)
    datos = {
        "personas": personas(wb),
        "anexo": anexo(wb),
        "memos_annie": memos_de_annie(wb),
    }
    for k, v in datos.items():
        print(f"{k}: {len(v)}", file=sys.stderr)
    json.dump(datos, sys.stdout, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
